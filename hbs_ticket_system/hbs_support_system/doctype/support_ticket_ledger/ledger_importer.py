import re
# Copyright (c) 2026, HBS and contributors
# For license information, please see license.txt

import os
import io
import json
import uuid
import datetime
import openpyxl
import xlrd
import frappe
from frappe import _


def normalize_phone(phone):
    """Normalize phone numbers to +91-XXXXXXXXXX format if valid."""
    if not phone:
        return ""
    phone = str(phone).strip()
    digits = "".join(filter(str.isdigit, phone))
    if len(digits) == 10:
        return f"+91-{digits}"
    elif len(digits) == 12 and digits.startswith("91"):
        return f"+91-{digits[2:]}"
    elif not phone.startswith("+") and digits:
        return f"+91-{digits}"
    return phone


def parse_date(val):
    """Convert various date formats or datetime objects to YYYY-MM-DD."""
    if not val:
        return None
    if isinstance(val, (datetime.date, datetime.datetime)):
        return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    if not s or s.lower() in ("none", "null", "-"):
        return None
    try:
        dt = frappe.utils.get_datetime(s)
        return dt.strftime("%Y-%m-%d")
    except Exception:
        return None


def get_user_match_map():
    """Build a lookup map for matching sales/installation executives to active Frappe users."""
    users = frappe.get_all(
        "User",
        filters={"enabled": 1},
        fields=["name", "full_name", "first_name", "email"]
    )
    user_map = {}
    for u in users:
        email = (u.email or u.name or "").strip().lower()
        full_name = (u.full_name or "").strip().lower()
        first_name = (u.first_name or "").strip().lower()

        if email:
            user_map[email] = u.name
        if full_name:
            user_map[full_name] = u.name
        if first_name and first_name not in user_map:
            user_map[first_name] = u.name

    return user_map


def match_user(name_str, user_map):
    """Attempt fuzzy/exact matching on user full name or first name."""
    if not name_str:
        return None
    clean = str(name_str).strip().lower()
    if not clean:
        return None
    if clean in user_map:
        return user_map[clean]
    for k, uid in user_map.items():
        if k and (clean == k or clean in k or k in clean):
            return uid
    return None


def get_or_create_contract_product(raw_type):
    """Resolve or create support_contract_product based on raw Type column."""
    raw = (raw_type or "").strip()
    if not raw or raw == "1":
        raw = "AMC"

    clean_upper = raw.upper()
    if clean_upper in ("AMC", "TALLY AMC", "ANNUAL MAINTENANCE CONTRACT"):
        category = "AMC"
        product_name = "Annual Maintenance Contract"
    elif "AMC" in clean_upper:
        category = "AMC"
        product_name = raw
    elif clean_upper in ("CLOUD", "TALLY ON CLOUD", "CLOUD RENEWAL"):
        category = "Cloud"
        product_name = "Tally on Cloud"
    elif "CLOUD" in clean_upper:
        category = "Cloud"
        product_name = raw
    elif clean_upper in ("MODULE", "MODULES", "CUSTOM MODULE"):
        category = "Module"
        product_name = "Custom Module"
    elif "MODUL" in clean_upper:
        category = "Module"
        product_name = raw
    else:
        category = "Other"
        product_name = raw

    # Ensure product exists in DB so Link validation passes
    if not frappe.db.exists("support_contract_product", product_name):
        p_doc = frappe.get_doc({
            "doctype": "support_contract_product",
            "product_name": product_name,
            "category": category,
            "default_period": "YEARLY",
            "standard_rate": 0.0
        })
        p_doc.insert(ignore_permissions=True)

    return product_name, category


def read_excel_rows(file_path):
    """Read .xlsx or .xls file robustly, checking magic bytes for modern zip-based vs legacy binary."""
    with open(file_path, "rb") as f:
        file_bytes = f.read()

    rows = []
    # Check if modern zip-based XLSX (PK magic bytes 0x50 0x4B)
    is_zip = file_bytes.startswith(b"PK")

    if is_zip:
        stream = io.BytesIO(file_bytes)
        wb = openpyxl.load_workbook(stream, data_only=True)
        sheet = wb.active
        if sheet.max_row > 1:
            raw_headers = [sheet.cell(1, c).value for c in range(1, sheet.max_column + 1)]
            headers = [str(h).strip() if h is not None else "" for h in raw_headers]
            for r in range(2, sheet.max_row + 1):
                row_dict = {}
                has_content = False
                for c in range(1, sheet.max_column + 1):
                    h = headers[c - 1] if c - 1 < len(headers) else f"col_{c}"
                    if not h:
                        continue
                    val = sheet.cell(r, c).value
                    if val is not None and str(val).strip():
                        has_content = True
                    row_dict[h] = val
                if has_content:
                    rows.append(row_dict)
    else:
        # Legacy binary XLS format
        wb = xlrd.open_workbook(file_contents=file_bytes)
        sheet = wb.sheet_by_index(0)
        if sheet.nrows > 1:
            raw_headers = sheet.row_values(0)
            headers = [str(h).strip() if h is not None else "" for h in raw_headers]
            for r in range(1, sheet.nrows):
                vals = sheet.row_values(r)
                row_dict = {}
                has_content = False
                for c, val in enumerate(vals):
                    h = headers[c] if c < len(headers) else f"col_{c}"
                    if not h:
                        continue
                    if val is not None and str(val).strip():
                        has_content = True
                    row_dict[h] = val
                if has_content:
                    rows.append(row_dict)

    return rows


def check_import_permission():
    user = frappe.session.user
    from hbs_ticket_system.hbs_support_system.doctype.support_ticket.support_ticket import is_owner_or_admin
    if not is_owner_or_admin(user):
        frappe.throw(
            _("Only Owners and Administrators are permitted to import Excel into Support Ticket Ledger."),
            frappe.PermissionError
        )


@frappe.whitelist()
def start_ledger_import(file_url):
    """Alias for prepare_import for backwards-compatibility."""
    return prepare_import(file_url)


@frappe.whitelist()
def prepare_import(file_url):
    """Uploads/locates file, extracts rows, caches them in Frappe cache, and returns summary preview."""
    check_import_permission()
    if not file_url:
        frappe.throw(_("Please upload or select an Excel file."))

    file_path = None
    if file_url.startswith("/files/"):
        file_path = frappe.get_site_path("public", file_url.lstrip("/"))
    elif file_url.startswith("/private/files/"):
        file_path = frappe.get_site_path("private", file_url.replace("/private/", ""))
    elif os.path.exists(file_url):
        file_path = file_url
    else:
        file_doc = frappe.db.get_value("File", {"file_url": file_url}, ["file_name"], as_dict=True)
        if file_doc:
            file_path = frappe.get_site_path("public", "files", file_doc.file_name)

    if not file_path or not os.path.exists(file_path):
        frappe.throw(_("Uploaded file could not be found at {0}").format(file_path or file_url))

    rows = read_excel_rows(file_path)
    if not rows:
        frappe.throw(_("No data rows found in the uploaded Excel file."))

    batch_token = str(uuid.uuid4())
    cache_key = f"ledger_import_{batch_token}"
    frappe.cache().set_value(cache_key, json.dumps(rows, default=str), expires_in_sec=3600)

    # Preview of first 3 rows
    preview = []
    for r in rows[:3]:
        s_no = r.get("Serial No") or r.get("Serial_No") or r.get("serial_no") or ""
        p_name = r.get("Party Name") or r.get("Party_Name") or r.get("party_name") or ""
        amt = r.get("Total Amount") or r.get("Billing Amount") or 0
        preview.append({
            "serial_no": str(s_no),
            "party_name": str(p_name),
            "amount": amt
        })

    return {
        "batch_token": batch_token,
        "total_rows": len(rows),
        "preview": preview
    }


@frappe.whitelist()
def process_import_batch(batch_token, start_idx=0, batch_size=50):
    """Process a chunk of rows from the cached import list."""
    check_import_permission()
    cache_key = f"ledger_import_{batch_token}"
    cached_data = frappe.cache().get_value(cache_key)
    if not cached_data:
        frappe.throw(_("Import session expired or not found. Please upload the file again."))

    rows = json.loads(cached_data)
    total_rows = len(rows)
    start_idx = int(start_idx or 0)
    batch_size = int(batch_size or 50)
    end_idx = min(start_idx + batch_size, total_rows)

    chunk = rows[start_idx:end_idx]
    user_map = get_user_match_map()

    inserted_count = 0
    updated_count = 0
    errors = []

    valid_companies = [
        "HBS Solutions Pvt Ltd",
        "NewAge InfoCom Pvt Ltd",
        "HBS Services"
    ]

    for idx, row in enumerate(chunk, start=start_idx + 1):
        try:
            # 1. Serial No (supports multiple comma-separated serials)
            raw_serial = row.get("Serial No") or row.get("Serial_No") or row.get("serial_no") or ""
            serial_candidates = [s.strip() for s in re.split(r"[,;/]+", str(raw_serial)) if s.strip()]
            valid_serials = []
            for s in serial_candidates:
                if s.endswith(".0"):
                    s = s[:-2]
                if s.isdigit() and len(s) == 9:
                    valid_serials.append(s)

            if not valid_serials:
                errors.append({
                    "row": idx,
                    "serial_no": str(raw_serial) or "Empty",
                    "party_name": str(row.get("Party Name") or "-"),
                    "reason": f"Invalid Serial No '{raw_serial}'. Must contain at least one valid 9-digit serial number."
                })
                continue
            serial_no = valid_serials[0]

            # 2. Party Name
            party_name = str(row.get("Party Name") or "").strip()
            if not party_name:
                errors.append({
                    "row": idx,
                    "serial_no": serial_no,
                    "party_name": "Empty",
                    "reason": "Party Name is required."
                })
                continue

            # 3. Company Name
            raw_comp = str(row.get("Company Name") or row.get("Our Company Name") or "").strip()
            our_company_name = "HBS Solutions Pvt Ltd"
            for vc in valid_companies:
                if raw_comp and vc.lower() in raw_comp.lower():
                    our_company_name = vc
                    break

            # 4. Email & Address
            email = str(row.get("Email") or "").strip()
            address = str(row.get("Address") or "").strip()
            pincode = str(row.get("Pincode") or "").strip()
            if pincode and pincode not in address:
                address = f"{address} {pincode}".strip()

            # 5. Dates & AMC Status
            start_date = parse_date(row.get("Starts From") or row.get("Start Date"))
            expiry_date = parse_date(row.get("Expiry Date"))
            status_val = str(row.get("Status") or "Active").strip()

            amc_status = "NO"
            if status_val.lower() == "active":
                amc_status = "YES"
            elif expiry_date:
                today_str = frappe.utils.nowdate()
                if expiry_date >= today_str:
                    amc_status = "YES"

            # 6. Billing Cycle / Period
            cycle = str(row.get("Billing Cycle") or "Yearly").strip().upper()
            if cycle not in ("YEARLY", "HALFYEARLY", "QUARTERLY", "MONTHLY"):
                cycle = "YEARLY"

            # 7. Amount & Visits
            amount = 0.0
            try:
                raw_amt = row.get("Total Amount") if row.get("Total Amount") is not None else row.get("Billing Amount")
                amount = float(raw_amt or 0.0)
            except Exception:
                amount = 0.0

            raw_visits = row.get("Visits")
            visits = str(raw_visits).strip() if raw_visits is not None else None

            # 8. Executives
            sold_by = row.get("Sold By")
            installed_by = row.get("Installed By")
            exec1 = match_user(sold_by, user_map)
            exec2 = match_user(installed_by, user_map)

            # 9. Contact Persons
            contact_person = str(row.get("Conctact Person") or row.get("Contact Person") or "").strip()
            phone = normalize_phone(row.get("Phone") or row.get("Phone Number") or row.get("Phone No"))
            mobile = normalize_phone(row.get("Mobile") or row.get("Mobile No") or row.get("Mobile Number"))

            contacts = []
            if contact_person or mobile or phone:
                contacts.append({
                    "contact_name": contact_person or "Primary Contact",
                    "contact_phone": mobile or "",
                    "phone_number": phone or ""
                })

            # 10. Multi-item Contract Details
            raw_type = (row.get("Type") or row.get("type") or row.get("Contract Type") or
                        row.get("Contract Name") or row.get("Product") or row.get("Category") or "").strip()
            product, category = get_or_create_contract_product(raw_type)

            raw_ip = str(row.get("IP Address") or "").strip()
            ip_address = raw_ip if raw_ip and raw_ip.lower() not in ("none", "null", "-") else None

            # For AMC, qty and rate remain empty (None); amount is direct
            is_amc = "AMC" in (category or "").upper()
            if is_amc:
                qty_val = None
                rate_val = None
            else:
                raw_qty = row.get("Users") or row.get("Qty")
                try:
                    qty_val = float(raw_qty) if raw_qty not in (None, "") else 1.0
                except Exception:
                    qty_val = 1.0
                rate_val = round(amount / qty_val, 2) if qty_val else amount

            contract_row = {
                "product": product,
                "category": category,
                "period": cycle,
                "start_date": start_date,
                "expiry_date": expiry_date,
                "visits": visits,
                "qty": qty_val,
                "rate": rate_val,
                "amount": amount,
                "ip_address": ip_address,
                "executive1": exec1,
                "executive2": exec2
            }

            # Match existing customer by any of the serial numbers
            matched_parent = None
            for s in valid_serials:
                p = frappe.db.get_value("support_ticket_ledger", {"serial_no": s}, "name")
                if not p:
                    p = frappe.db.get_value("support_ticket_ledger", s, "name")
                if not p:
                    p = frappe.db.get_value("support_ticket_ledger_serials", {"serial_no": s}, "parent")
                if p:
                    matched_parent = p
                    break

            if matched_parent and frappe.db.exists("support_ticket_ledger", matched_parent):
                doc = frappe.get_doc("support_ticket_ledger", matched_parent)
                doc.party_name = party_name
                if email:
                    doc.email = email
                doc.our_company_name = our_company_name
                if address:
                    doc.address = address
                doc.amc_status = amc_status
                doc.amc_period = cycle
                if start_date:
                    doc.amc_start_date = start_date
                if expiry_date:
                    doc.amc_expiry_date = expiry_date
                doc.amount = amount
                if visits:
                    doc.visits = visits
                if exec1:
                    doc.executive = exec1
                if exec2:
                    doc.executive2 = exec2

                # Sync all valid serial numbers into child table
                existing_serials = {s_row.serial_no for s_row in (doc.get("serial_numbers") or []) if s_row.serial_no}
                for s in valid_serials:
                    if s not in existing_serials:
                        doc.append("serial_numbers", {"serial_no": s})
                        existing_serials.add(s)

                # Update / Append contact details
                if contacts:
                    new_c = contacts[0]
                    # Remove legacy "(Office)" child rows for this contact
                    for off_r in [r for r in (doc.get("contact_details") or []) if (r.contact_name or "").strip().endswith("(Office)")]:
                        if not phone and off_r.contact_phone:
                            new_c["phone_number"] = off_r.contact_phone
                        doc.remove(off_r)

                    if doc.get("contact_details"):
                        c_row = doc.contact_details[0]
                        if contact_person:
                            c_row.contact_name = contact_person
                        if mobile:
                            c_row.contact_phone = mobile
                        if phone or new_c.get("phone_number"):
                            c_row.phone_number = phone or new_c.get("phone_number")
                    else:
                        doc.append("contact_details", new_c)

                # Sync contract row into contracts child table
                # Preserves history: Only updates if exact product AND same period (re-upload of same file).
                # New renewals with new dates or different AMC products are appended as separate rows.
                if frappe.db.exists("DocType", "support_ticket_ledger_contracts"):
                    matched_contract = None
                    for c_row in (doc.get("contracts") or []):
                        same_product = (c_row.product == product)
                        same_period = False
                        if start_date and c_row.start_date and str(c_row.start_date) == str(start_date):
                            same_period = True
                        elif expiry_date and c_row.expiry_date and str(c_row.expiry_date) == str(expiry_date):
                            same_period = True
                        elif not start_date and not expiry_date:
                            same_period = True

                        if same_product and same_period:
                            matched_contract = c_row
                            break

                    if matched_contract:
                        matched_contract.product = product
                        matched_contract.category = category
                        matched_contract.period = cycle
                        if start_date:
                            matched_contract.start_date = start_date
                        if expiry_date:
                            matched_contract.expiry_date = expiry_date
                        matched_contract.visits = visits
                        matched_contract.qty = qty_val
                        matched_contract.rate = rate_val
                        matched_contract.amount = amount
                        if ip_address:
                            matched_contract.ip_address = ip_address
                        if exec1:
                            matched_contract.executive1 = exec1
                        if exec2:
                            matched_contract.executive2 = exec2
                    else:
                        doc.append("contracts", contract_row)

                doc.save(ignore_permissions=True)
                updated_count += 1
            else:
                serial_rows = [{"serial_no": s} for s in valid_serials]
                doc_dict = {
                    "doctype": "support_ticket_ledger",
                    "party_name": party_name,
                    "serial_no": valid_serials[0],
                    "email": email or None,
                    "our_company_name": our_company_name,
                    "address": address or None,
                    "amc_status": amc_status,
                    "amc_period": cycle,
                    "amc_start_date": start_date,
                    "amc_expiry_date": expiry_date,
                    "amount": amount,
                    "visits": visits,
                    "executive": exec1,
                    "executive2": exec2,
                    "serial_numbers": serial_rows,
                    "contact_details": contacts
                }
                if frappe.db.exists("DocType", "support_ticket_ledger_contracts"):
                    doc_dict["contracts"] = [contract_row]

                new_doc = frappe.get_doc(doc_dict)
                new_doc.insert(ignore_permissions=True)
                inserted_count += 1

        except Exception as e:
            frappe.log_error(f"Error importing row {idx}: {str(e)}", "Ledger Excel Import")
            errors.append({
                "row": idx,
                "serial_no": str(row.get("Serial No") or "-"),
                "party_name": str(row.get("Party Name") or "-"),
                "reason": str(e)
            })

    frappe.db.commit()

    is_complete = end_idx >= total_rows
    progress_pct = int((end_idx / total_rows) * 100) if total_rows else 100

    return {
        "start_idx": start_idx,
        "processed_count": end_idx - start_idx,
        "end_idx": end_idx,
        "total_rows": total_rows,
        "inserted": inserted_count,
        "updated": updated_count,
        "errors": errors,
        "is_complete": is_complete,
        "progress_pct": progress_pct
    }
